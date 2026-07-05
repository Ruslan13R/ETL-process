from openpyxl import load_workbook
import logging

class Extractor:
    def exctract(self):
        wd = load_workbook('./data/kirby-msk.xlsx')
        ws = wd.active

        list_of_rows = []

        for row in ws.iter_rows(values_only=True):
            list_of_rows.append(row)

        headers = list_of_rows.pop(0)

        return headers, list_of_rows
